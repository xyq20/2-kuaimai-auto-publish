"""Local workbench: trusted product catalog and durable, serial CLI jobs."""
from __future__ import annotations

from contextlib import contextmanager
from datetime import datetime
import fcntl
import hashlib
import json
import os
from pathlib import Path
import re
import signal
import subprocess
import sys
import threading
import uuid
from zipfile import BadZipFile

from .config import PROJECT_DIR, Settings
from .database import connect, utc_now
from .service import ApiError, require
from platform_registry import PLATFORM_SPECS

PLATFORMS = {s.cli_name: s.display_name.replace("资料", "") for s in PLATFORM_SPECS if s.enabled_in_all}
ACTIVE = {"running", "waiting_review", "stopping"}
IDENTIFIER = re.compile(r"^[0-9a-f]{32}$")


def atomic_json(path: Path, data: dict) -> None:
    temporary = path.with_suffix(".tmp")
    temporary.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def read_json(path: Path, default=None):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return default


class ProductCatalog:
    def __init__(self, root: Path, runs_root: Path = None):
        self.root = root.resolve()
        self.runs_root = runs_root
        self._cache = {}

    def _safe(self, path: Path) -> bool:
        return path.resolve().is_relative_to(self.root)

    def paths(self) -> dict:
        if not self.root.is_dir():
            raise ApiError(503, "products_unavailable")
        try:
            workbooks = {}
            for path in sorted(self.root.glob("*/*")):
                if (
                    re.sub(r"\s+", "", path.name).casefold() == "产品信息.xlsx"
                    and self._safe(path) and path.is_file()
                ):
                    # One product per folder; prefer the canonical workbook
                    # when both a normal name and a whitespace variant exist.
                    if path.parent not in workbooks or path.name == "产品信息.xlsx":
                        workbooks[path.parent] = path
            return {
                hashlib.sha256(str(p.resolve()).encode()).hexdigest()[:32]: p
                for p in workbooks.values()
            }
        except OSError:
            raise ApiError(503, "products_unavailable") from None

    def path(self, product_id: str) -> Path:
        require(isinstance(product_id, str) and IDENTIFIER.fullmatch(product_id), 404, "product_not_found")
        path = self.paths().get(product_id)
        require(path is not None, 404, "product_not_found")
        return path

    def listing(self) -> list:
        paths = self.paths()
        progress = {p.parent.resolve(): set() for p in paths.values()}
        if self.runs_root is not None:
            for summary_path in self.runs_root.glob("*/input-summary.json"):
                summary = read_json(summary_path, {})
                excel_path = summary.get("excel_path")
                if not excel_path:
                    continue
                folder = Path(excel_path).parent.resolve()
                if folder not in progress:
                    continue
                run = summary_path.parent
                receipts = [*run.glob("save-result.json"), *run.glob("publish-result.json"),
                            *run.glob("*/save-result.json"), *run.glob("*/publish-result.json")]
                for receipt_path in receipts:
                    receipt = read_json(receipt_path, {})
                    if receipt.get("result") != 1:
                        continue
                    platform = receipt_path.parent.name
                    label = PLATFORMS.get(platform)
                    if label is None:
                        label = receipt.get("shop_publish", {}).get("platform")
                    if label is None:
                        validations = list(receipt_path.parent.glob("*-after-save-validation.json"))
                        label = next((PLATFORMS.get(p.name.split('-after-save')[0])
                                      for p in validations if p.name.split('-after-save')[0] in PLATFORMS), None)
                    progress[folder].add(label or "商品资料")
        return [{"id": key, "title": p.parent.name,
                 "processed": bool(progress[p.parent.resolve()]),
                 "processed_platforms": sorted(progress[p.parent.resolve()])}
                for key, p in paths.items()]

    def detail(self, product_id: str) -> dict:
        path = self.path(product_id)
        try:
            signature = (path.stat().st_mtime_ns, path.stat().st_size)
            cached = self._cache.get(product_id)
            if cached and cached[0] == signature:
                return cached[1]
            from openpyxl import load_workbook
            from kuaimai_erp import read_excel_fields, normalize_cell
            from douyin_data import field_lookup
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                sheet = workbook[workbook.sheetnames[0]]
                fields = read_excel_fields(list(sheet.iter_rows(values_only=True)))
            finally:
                workbook.close()
            def value(*aliases):
                match = field_lookup(fields, *aliases)
                return normalize_cell(match[1]) if match else ""
            detail = {
                "id": product_id,
                "title": value("商品标题", "商品名称", "宝贝标题") or path.parent.name,
                "style_code": value("货号", "商家外部编码", "款式编码") or path.parent.name,
                "category": value("商品分类") or "以 Excel 类目为准",
                "folder": path.parent.name,
                "image": f"/api/launcher/products/{product_id}/image" if self.image(product_id) else None,
            }
            self._cache[product_id] = (signature, detail)
            return detail
        except (OSError, ValueError, KeyError, BadZipFile):
            raise ApiError(422, "product_unreadable") from None

    def image(self, product_id: str):
        folder = self.path(product_id).parent
        for name in ("1：1主图", "1:1主图", "1：1 主图", "主图"):
            directory = folder / name
            if not self._safe(directory) or not directory.is_dir():
                continue
            for path in sorted(directory.iterdir()):
                if path.suffix.lower() in {".jpg", ".jpeg", ".png", ".webp"} and self._safe(path) and path.is_file():
                    return path
        return None


def build_command(product: Path, config: dict, run_id: str, output_dir: Path, settings: Settings) -> list:
    operation = config.get("operation", "platforms")
    require(operation in {"platforms", "base", "create"}, 422, "invalid_operation")
    mode = config.get("mode")
    require(mode in {"preview", "save", "publish"}, 422, "invalid_mode")
    names = config.get("platforms")
    if operation == "platforms":
        require(isinstance(names, list) and 0 < len(names) <= len(PLATFORMS), 422, "invalid_platforms")
        require(all(isinstance(n, str) and n in PLATFORMS for n in names), 422, "invalid_platforms")
        require(len(set(names)) == len(names), 422, "invalid_platforms")
    require(type(config.get("learning")) is bool, 422, "invalid_learning")
    # Base editing currently saves as part of its workflow. Never expose it as a preview.
    require(operation == "platforms" or mode == "save", 422, "base_requires_save")
    require(mode != "publish" or config.get("confirm_publish") is True, 422, "publish_confirmation_required")
    learning = config["learning"] and operation == "platforms"
    require(not learning or settings.device_token, 503, "review_not_configured")
    command = [sys.executable, "-u", str(PROJECT_DIR / "kuaimai_erp.py"),
               "--excel-url", str(product), "--platform", ",".join(names) if operation == "platforms" else "base",
               "--run-id", run_id, "--output-dir", str(output_dir),
               {"preview": "--no-save", "save": "--save-only", "publish": "--save"}[mode]]
    if operation == "create":
        command.append("--create-product")
    if names == ["taobao"] and operation == "platforms" and mode != "preview":
        command.append("--allow-taobao-save-once" if mode == "save" else "--allow-taobao-publish-once")
    if learning:
        command += ["--learning-enabled", "--learning-api-url", f"http://127.0.0.1:{settings.review_port}"]
    return command


def process_command(pid: int) -> str:
    try:
        return subprocess.check_output(["ps", "-ww", "-p", str(pid), "-o", "command="], text=True, errors="replace", timeout=3, stderr=subprocess.DEVNULL).strip()
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired, OSError):
        return ""


def external_runner_active() -> bool:
    """Also respect a runner started from the existing command-line launcher."""
    try:
        lines = subprocess.check_output(["ps", "-ww", "-Ao", "comm=,args="], text=True, errors="replace", timeout=3).splitlines()
        return any(re.match(r"\S*(?:python|Python)\S*\s", line.strip()) and
                   re.search(r"(?:^|[/\s])kuaimai_erp\.py(?:\s|$)", line) for line in lines)
    except (OSError, subprocess.CalledProcessError, subprocess.TimeoutExpired):
        raise ApiError(503, "process_check_failed") from None


class Launcher:
    def __init__(self, settings: Settings):
        self.settings = settings
        self.catalog = ProductCatalog(settings.products_root, PROJECT_DIR / "output/kuaimai/runs")
        self.root = settings.data_dir.resolve() / "launcher"
        self.root.mkdir(parents=True, exist_ok=True)
        self._mutex = threading.Lock()
        self._children = {}

    @contextmanager
    def locked(self):
        with self._mutex, (self.root / ".lock").open("a") as handle:
            fcntl.flock(handle, fcntl.LOCK_EX)
            try:
                yield
            finally:
                fcntl.flock(handle, fcntl.LOCK_UN)

    def _directory(self, job_id: str) -> Path:
        require(isinstance(job_id, str) and IDENTIFIER.fullmatch(job_id), 404, "job_not_found")
        return self.root / job_id

    def _owned(self, job: dict) -> bool:
        child = self._children.get(job["id"])
        if child is not None and child.poll() is not None:
            self._children.pop(job["id"], None)
            return False
        command = process_command(job.get("pid", 0))
        return "local_review.launcher_worker" in command and str(self.root / job["id"]) in command

    def _refresh(self, job: dict) -> dict:
        child = self._children.get(job["id"])
        if child is not None and child.poll() is not None:
            self._children.pop(job["id"], None)
        if job["status"] in ACTIVE:
            receipt = read_json(self.root / job["id"] / "result.json")
            if receipt is not None:
                job.update(status="stopped" if receipt.get("cancelled") or job.get("stop_requested") else
                           "completed" if receipt["exit_code"] == 0 else "failed",
                           exit_code=receipt["exit_code"], finished_at=receipt["finished_at"])
            elif not self._owned(job):
                job.update(status="stopped" if job.get("stop_requested") else "interrupted", finished_at=utc_now())
            if job["status"] in {"stopped", "interrupted", "failed"}:
                # A dead runner cannot consume a review. Hide its stale queue using
                # the same terminal checkpoint rule as the existing review center.
                with connect(self.settings) as connection:
                    connection.execute(
                        "UPDATE run_checkpoints SET status='failed',version=version+1,updated_at=? "
                        "WHERE run_id=? AND status NOT IN ('completed','failed','cancelled')",
                        (utc_now(), job["id"]),
                    )
            atomic_json(self.root / job["id"] / "job.json", job)
        return job

    def _jobs(self) -> list:
        return sorted([self._refresh(job) for path in self.root.glob("*/job.json")
                       if (job := read_json(path))], key=lambda j: j["created_at"], reverse=True)

    def _public(self, job: dict, *, logs=False) -> dict:
        result = {k: v for k, v in job.items() if k not in {"pid", "output_dir", "request_key"}}
        stages = {name: "pending" for name in job["platforms"]}
        receipts = read_json(Path(job["output_dir"]) / "all-platform-result.json", [])
        for item in receipts:
            if item.get("platform") in stages:
                stages[item["platform"]] = item["status"]
        with connect(self.settings) as connection:
            checkpoint = connection.execute("SELECT current_index,status FROM run_checkpoints WHERE run_id=?", (job["id"],)).fetchone()
            pending = connection.execute("SELECT COUNT(*) FROM review_tasks WHERE run_id=? AND status IN ('pending','claimed')", (job["id"],)).fetchone()[0]
        current = checkpoint["current_index"] if checkpoint else next((i for i, s in enumerate(stages.values()) if s != "success"), len(stages))
        if checkpoint:
            for name in list(stages)[:current]:
                stages[name] = "success"
        if job["status"] == "completed":
            stages = {name: "success" for name in stages}
        elif job["status"] in ACTIVE and current < len(stages):
            waiting = pending > 0 and checkpoint and checkpoint["status"] == "waiting_review"
            result["status"] = "stopping" if job["status"] == "stopping" else "waiting_review" if waiting else "running"
            stages[list(stages)[current]] = result["status"]
        elif current < len(stages):
            stages[list(stages)[current]] = job["status"]
        result.update(stages=[{"platform": name, "status": status} for name, status in stages.items()], pending_reviews=pending)
        if logs:
            try:
                with (self.root / job["id"] / "run.log").open("rb") as handle:
                    handle.seek(0, 2)
                    handle.seek(max(0, handle.tell() - 48000))
                    result["log"] = handle.read().decode("utf-8", errors="replace")
            except OSError:
                result["log"] = ""
            for secret in (self.settings.device_token, self.settings.model_api_key):
                if secret:
                    result["log"] = result["log"].replace(secret, "[已隐藏]")
        return result

    def history(self) -> list:
        with self.locked():
            return [self._public(job) for job in self._jobs()[:100]]

    def get(self, job_id: str) -> dict:
        with self.locked():
            job = read_json(self._directory(job_id) / "job.json")
            require(job is not None, 404, "job_not_found")
            return self._public(self._refresh(job), logs=True)

    def start(self, payload: dict, username: str) -> dict:
        key = payload.get("request_key")
        require(isinstance(key, str) and IDENTIFIER.fullmatch(key), 422, "invalid_request_key")
        with self.locked():
            jobs = self._jobs()
            existing = next((job for job in jobs if job["request_key"] == key), None)
            if existing:
                return self._public(existing)
            require(not any(job["status"] in ACTIVE for job in jobs), 409, "job_already_running")
            require(not external_runner_active(), 409, "external_job_running")
            product = self.catalog.path(payload.get("product_id"))
            detail = self.catalog.detail(payload["product_id"])
            job_id = uuid.uuid4().hex
            output_dir = PROJECT_DIR / "output/kuaimai/runs" / (datetime.now().strftime("%Y%m%d-%H%M%S") + "-" + job_id[:8])
            command = build_command(product, payload, job_id, output_dir, self.settings)
            directory = self.root / job_id
            directory.mkdir()
            atomic_json(directory / "launch.json", {"argv": command})
            job = {"id": job_id, "request_key": key, "product_id": detail["id"], "title": detail["title"],
                   "style_code": detail["style_code"], "platforms": payload["platforms"] if payload.get("operation", "platforms") == "platforms" else ["base"],
                   "mode": payload["mode"], "operation": payload.get("operation", "platforms"),
                   "learning": "--learning-enabled" in command, "status": "running", "created_at": utc_now(),
                   "created_by": username, "output_dir": str(output_dir)}
            env = os.environ.copy()
            env["KUAIMAI_LEARNING_DEVICE_TOKEN"] = self.settings.device_token
            env["PYTHONUNBUFFERED"] = "1"
            atomic_json(directory / "job.json", job)
            try:
                with (directory / "run.log").open("ab") as log:
                    child = subprocess.Popen([sys.executable, "-u", "-m", "local_review.launcher_worker", str(directory)],
                                             cwd=PROJECT_DIR, env=env, stdin=subprocess.DEVNULL, stdout=log, stderr=subprocess.STDOUT,
                                             start_new_session=True)
                self._children[job_id] = child
                job["pid"] = child.pid
            except OSError:
                job.update(status="failed", finished_at=utc_now(), exit_code=-1)
                atomic_json(directory / "job.json", job)
                raise ApiError(503, "launch_failed") from None
            atomic_json(directory / "job.json", job)
            return self._public(job)

    def stop(self, job_id: str) -> dict:
        with self.locked():
            directory = self._directory(job_id)
            job = read_json(directory / "job.json")
            require(job is not None, 404, "job_not_found")
            self._refresh(job)
            if job["status"] in ACTIVE and self._owned(job):
                job.update(status="stopping", stop_requested=True)
                atomic_json(directory / "job.json", job)
                try:
                    os.killpg(job["pid"], signal.SIGINT)
                except ProcessLookupError:
                    pass
            return self._public(job)
